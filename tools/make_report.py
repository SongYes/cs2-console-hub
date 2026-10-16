# -*- coding: utf-8 -*-
"""生成「翻译修复前后对照」报告。

产出：
  docs/translation-report.docx / .pdf   交给你审阅的对照文档
  docs/translation-report.html          完整对照表（可搜索，含全部条目）
  docs/translation-report.json          机器可读的完整对照数据

「修复前」指的是上一次构建里实际展示给用户的中文释义：
  - 中英混杂的规则翻译（例如「默认值: 1 设置 the alternate 雷达 缩放. Valid 取值 are 0.25 to 1.0」）
  - 只有英文原文、没有中文的条目
「修复后」是本次机器翻译 + 术语规范化的结果。
"""
import html
import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
DOCS = ROOT / "docs"
DOCS.mkdir(exist_ok=True)

BAD_ZH_LIMIT = 120   # 报告里「修复前」最多展示的字符数

CAT_NAME = {}


def en_words(s):
    return len(re.findall(r"[A-Za-z]{2,}", s or ""))


def cn_chars(s):
    return len(re.findall(r"[\u4e00-\u9fff]", s or ""))


def looks_bad(zh):
    """判定「中文翻译不好」：仍以英文为主、或中英混杂。"""
    if not zh:
        return False
    en, cn = en_words(zh), cn_chars(zh)
    if en >= 2:
        return True
    if en >= 1 and cn < 8:
        return True
    return False


def load():
    db = json.loads((DATA / "commands.json").read_text(encoding="utf-8"))
    mt = json.loads((DATA / "mt_cache.json").read_text(encoding="utf-8"))
    mt = {k: v["zh"] for k, v in mt.items() if isinstance(v, dict) and v.get("zh")}
    prev_path = DATA / "zh_prev.json"
    prev = json.loads(prev_path.read_text(encoding="utf-8")) if prev_path.exists() else {}
    for c in db["categories"]:
        CAT_NAME[c["id"]] = c["name"]
    return db, mt, prev


SRC_NAME = {"mymemory": "在线引擎", "argos": "离线模型", "mt": "机器翻译"}


def collect(db, mt, prev):
    """挑出所有「修复前不理想 / 修复后变好」的条目。"""
    rows = []
    for e in db["entries"]:
        new = e["z"] or ""
        old = e.get("z0") or prev.get(e["n"]) or ""
        if not e["en"]:
            continue
        changed = old != new
        bad_before = looks_bad(old) or not re.search(r"[\u4e00-\u9fff]", old or "")
        if not (changed and bad_before):
            continue
        rows.append({
            "name": e["n"], "cat": CAT_NAME.get(e["c"], e["c"]), "sub": e["u"],
            "en": e["en"], "old": old, "new": new,
            "q": e.get("q"), "k": e.get("k"), "r": e.get("r"),
            "ms": SRC_NAME.get(e.get("ms") or "", ""),
            "kind": ("精校" if e.get("q") == 2 else "机翻" if e.get("q") == 3
                     else "规则" if e.get("q") == 1 else "原文"),
            "fixed": changed,
        })
    # 展示顺序：先实用指令，再按分类、名称
    order = {c["id"]: i for i, c in enumerate(db["categories"])}
    rows.sort(key=lambda r: (0 if r["k"] else 1, order.get(r["cat"], 99), r["name"].lower()))
    return rows


def fmt_old(row):
    """把「修复前」渲染成一行可读文本。"""
    if row["old"] and looks_bad(row["old"]):
        return row["old"][:BAD_ZH_LIMIT] + ("…" if len(row["old"]) > BAD_ZH_LIMIT else "")
    return "（无中文，仅英文原文）"


def build_html(rows, stats, path):
    esc = html.escape
    head = """<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
<title>CS2 指令库 中文释义修复对照</title>
<style>
 body{font:14px/1.7 "Microsoft YaHei",system-ui,sans-serif;margin:0;background:#0f1319;color:#e8edf5}
 header{padding:28px 34px;border-bottom:1px solid #26303f;background:#151b24;position:sticky;top:0;z-index:5}
 h1{margin:0 0 6px;font-size:21px}
 .meta{color:#9aa7bb;font-size:13px}
 .stats{display:flex;flex-wrap:wrap;gap:10px;margin-top:14px}
 .stat{background:#1b2230;border:1px solid #2b3548;border-radius:9px;padding:8px 14px;font-size:13px}
 .stat b{color:#f0a92b;font-size:16px}
 main{padding:18px 34px 60px}
 table{border-collapse:collapse;width:100%;font-size:13px}
 th,td{border-bottom:1px solid #232c3a;padding:9px 10px;vertical-align:top;text-align:left}
 th{background:#1b2230;color:#9aa7bb;font-weight:600;position:sticky;top:96px}
 td.n{font-family:Consolas,monospace;color:#ffc75c;white-space:nowrap}
 td.en{color:#8d9bb0;font-family:Consolas,monospace;font-size:12.3px;max-width:330px}
 td.old{color:#ff9b9b;max-width:330px}
 td.new{color:#b9f5d0;max-width:380px}
 .tag{display:inline-block;font-size:11px;padding:1px 7px;border-radius:999px;border:1px solid #33405a;color:#9aa7bb;margin-right:5px}
 .tag.k{color:#45c98a;border-color:#2c6b4c}
 input{background:#1b2230;border:1px solid #33405a;color:#e8edf5;border-radius:7px;padding:8px 12px;width:320px;outline:none}
 input:focus{border-color:#f0a92b}
</style></head><body>
<header>
 <h1>CS2 控制台指令查询中心 · 中文释义修复对照</h1>
 <div class="meta">生成时间：__AT__ ｜ 数据源：__SOURCE__ ｜ 条目总数 __TOTAL__</div>
 <div class="stats">
   <div class="stat">修复条目 <b>__FIXED__</b></div>
   <div class="stat">修复前「仅英文」 <b>__OLDE__</b></div>
   <div class="stat">修复前「中英混杂」 <b>__OLDM__</b></div>
   <div class="stat">现在有中文释义 <b>__WITHZH__</b></div>
   <div class="stat">仍残留英文词 <b>__RESID__</b></div>
   <div class="stat">其中实用指令 <b>__CORE__</b></div>
 </div>
 <p class="meta" style="margin:14px 0 0"><input id="q" placeholder="过滤：指令名 / 英文 / 中文"></p>
</header><main><table id="t">
<thead><tr><th style="width:210px">指令</th><th style="width:90px">分类</th><th>修复前（旧中文释义）</th><th>修复后（新中文释义）</th><th style="width:300px">英文原文</th></tr></thead>
<tbody>
"""
    for k, v in (("__AT__", stats["at"]), ("__SOURCE__", esc(stats["source"])),
                 ("__TOTAL__", stats["total"]), ("__FIXED__", stats["fixed"]),
                 ("__OLDE__", stats["old_english"]), ("__OLDM__", stats["old_mixed"]),
                 ("__WITHZH__", stats["with_zh"]), ("__RESID__", stats["residual"]),
                 ("__CORE__", stats["core_fixed"])):
        head = head.replace(k, str(v))
    parts = [head]

    for r in rows:
        tags = f'<span class="tag">{"精校" if r["q"] == 2 else "机翻" if r["q"] == 3 else "规则" if r["q"] == 1 else "原文"}</span>'
        if r["k"]:
            tags += '<span class="tag k">实用</span>'
        parts.append(
            f'<tr data-s="{esc((r["name"] + " " + r["en"] + " " + r["new"]).lower())}">'
            f'<td class="n">{esc(r["name"])}<br>{tags}</td>'
            f'<td>{esc(r["cat"])}</td>'
            f'<td class="old">{esc(fmt_old(r))}</td>'
            f'<td class="new">{esc(r["new"])}</td>'
            f'<td class="en">{esc(r["en"][:300])}</td></tr>')

    parts.append("""</tbody></table></main>
<script>
document.getElementById('q').addEventListener('input', e => {
  const k = e.target.value.toLowerCase().trim();
  document.querySelectorAll('#t tbody tr').forEach(tr => {
    tr.style.display = (!k || tr.dataset.s.includes(k)) ? '' : 'none';
  });
});
</script></body></html>""")
    path.write_text("".join(parts), encoding="utf-8")
    return path


# --------------------------------------------------------------------------- #
# DOCX / PDF
# --------------------------------------------------------------------------- #

def set_cjk(run, ascii_font="Microsoft YaHei", ea_font="Microsoft YaHei", size=None, bold=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ea_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def build_docx(rows, stats, path, limit=None):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.4)

    # 正文默认字体
    st = doc.styles["Normal"]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(9)
    from docx.oxml.ns import qn
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    def para(text, size=9, bold=False, color=None, space_after=4, align=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        if align:
            p.alignment = align
        r = p.add_run(text)
        set_cjk(r, size=size, bold=bold)
        if color:
            r.font.color.rgb = RGBColor(*color)
        return p

    # ---- 封面/摘要 ----
    para("CS2 控制台指令查询中心", size=20, bold=True, color=(0x1F, 0x2A, 0x37), space_after=2)
    para("中文释义修复前后对照报告", size=14, bold=True, color=(0xB7, 0x74, 0x00), space_after=10)
    para(f"生成时间：{stats['at']}", size=9, color=(0x60, 0x6A, 0x78))
    para(f"指令库版本：{stats['version']}　条目总数：{stats['total']}", size=9, color=(0x60, 0x6A, 0x78))
    para(f"数据来源：{stats['source']}", size=9, color=(0x60, 0x6A, 0x78), space_after=12)

    para("一、处理说明", size=12, bold=True, color=(0x1F, 0x2A, 0x37), space_after=6)
    for line in [
        "1. 修复对象：指令库中「没有中文翻译」和「中文翻译不良好（中英混杂、英文残留过多）」的条目。",
        "2. 修复方式：对每条英文官方说明做整句中文翻译，再按 CS2 玩家习惯做术语规范化",
        "　 （例如 ragdoll → 尸体模型、观察者 → 观战者、X-Ray → X 光、HLTV / SourceTV 保留原名）。",
        "3. 优先级：随正式版下发的实用指令优先，其次是引擎内部变量。",
        "4. 未改动的部分：已人工精校的 %d 条释义保持原样。"
        % stats["hand"],
        "5. 仍含英文的部分：HLTV、SourceTV、Steam、UDP、SDR 等专有名词与协议名保留原文，属于有意为之。",
    ]:
        para(line, size=9, space_after=2)

    para("二、结果统计", size=12, bold=True, color=(0x1F, 0x2A, 0x37), space_after=6)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    for i, t in enumerate(("指标", "数值")):
        hdr[i].text = ""
        set_cjk(hdr[i].paragraphs[0].add_run(t), size=9, bold=True)
    for k, v in [
        ("条目总数", stats["total"]),
        ("本次修复条目", stats["fixed"]),
        ("其中实用指令（默认视图可见）", stats["core_fixed"]),
        ("修复前：仅英文原文、无中文", stats["old_english"]),
        ("修复前：中英混杂", stats["old_mixed"]),
        ("修复后：有中文释义", stats["with_zh"]),
        ("修复后：仍含专有名词英文", stats["residual"]),
        ("人工精校（未改动）", stats["hand"]),
    ]:
        cells = tbl.add_row().cells
        for c, t in zip(cells, (str(k), str(v))):
            c.text = ""
            set_cjk(c.paragraphs[0].add_run(t), size=9)

    doc.add_page_break()

    show = rows if limit is None else rows[:limit]
    note = "" if limit is None or limit >= len(rows) else \
        f"（为控制篇幅，此处只列前 {len(show)} 条；完整 {len(rows)} 条见同名 HTML / JSON 文件）"
    para(f"三、逐条对照（共 {len(rows)} 条）{note}",
         size=12, bold=True, color=(0x1F, 0x2A, 0x37), space_after=8)

    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.autofit = False
    widths = (Cm(3.6), Cm(4.6), Cm(5.4), Cm(4.4))
    head = t.rows[0].cells
    for i, (txt, w) in enumerate(zip(("指令 / 分类", "修复前（旧释义）", "修复后（新释义）", "英文原文"), widths)):
        head[i].text = ""
        set_cjk(head[i].paragraphs[0].add_run(txt), size=8.5, bold=True)
        head[i].width = w

    for r in show:
        cells = t.add_row().cells
        vals = (
            (f"{r['name']}\n（{r['cat']}）", None, 8),
            (fmt_old(r), (0xB0, 0x30, 0x30), 8),
            (r["new"], (0x1E, 0x6B, 0x45), 8.5),
            (r["en"][:220], (0x60, 0x6A, 0x78), 7.5),
        )
        for i, (txt, col, sz) in enumerate(vals):
            c = cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            run = p.add_run(txt)
            set_cjk(run, size=sz)
            if col:
                run.font.color.rgb = RGBColor(*col)
            c.width = widths[i]

    doc.save(path)
    return path


def to_pdf(docx_path, pdf_path):
    """用本机 Word 把 DOCX 导出为 PDF。

    先导出到临时文件再搬过来：Word 有时会误判目标 PDF「只读」，
    直接写目标路径会失败；走临时文件 + 重试能稳定绕开。
    """
    import os
    import shutil
    import time
    import win32com.client as win32          # noqa

    pdf_path = Path(pdf_path)
    tmp_pdf = pdf_path.with_name(pdf_path.stem + "_tmp.pdf")
    for p in (tmp_pdf, pdf_path):
        try:
            if p.exists():
                p.unlink()
        except Exception:
            pass

    last = None
    for attempt in range(3):
        word = win32.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            doc = word.Documents.Open(str(docx_path), ReadOnly=True, AddToRecentFiles=False)
            doc.ExportAsFixedFormat(str(tmp_pdf), 17)   # 17 = wdExportFormatPDF
            doc.Close(False)
        except Exception as e:                          # noqa
            last = e
            try:
                word.Quit()
            except Exception:
                pass
            time.sleep(2)
            continue
        finally:
            try:
                word.Quit()
            except Exception:
                pass
        if tmp_pdf.exists() and tmp_pdf.stat().st_size > 1000:
            try:
                os.replace(str(tmp_pdf), str(pdf_path))     # 原子替换，比 shutil.move 稳
                return pdf_path
            except Exception as e:                          # 极端情况：目标被占用则改名保存
                alt = pdf_path.with_name(f"{pdf_path.stem}-new{pdf_path.suffix}")
                try:
                    if alt.exists():
                        alt.unlink()
                    os.replace(str(tmp_pdf), str(alt))
                    print(f"[i] 目标 PDF 被占用，已输出为新文件：{alt.name}")
                    return alt
                except Exception:
                    last = e
        last = RuntimeError("Word 未产出文件")
        time.sleep(2)
    raise RuntimeError(f"导出 PDF 失败：{last}")


def main():
    # 可选参数：PDF 里最多列多少条（默认不限）；HTML/JSON 始终包含全部
    limit = None
    for a in sys.argv[1:]:
        if a.startswith("--pdf-limit="):
            limit = int(a.split("=", 1)[1])
        elif a.isdigit():
            limit = int(a)
    db, mt, prev = load()
    rows = collect(db, mt, prev)
    total = db["count"]

    fixed = [r for r in rows if r["fixed"]]
    old_english = sum(1 for r in fixed if not r["old"] or not re.search(r"[\u4e00-\u9fff]", r["old"]))
    old_mixed = sum(1 for r in fixed if r["old"] and re.search(r"[\u4e00-\u9fff]", r["old"])
                    and looks_bad(r["old"]))
    with_zh = sum(1 for e in db["entries"] if e["z"] and re.search(r"[\u4e00-\u9fff]", e["z"]))
    residual = sum(1 for r in fixed if en_words(r["new"]) >= 2)
    hand = sum(1 for e in db["entries"] if e.get("q") == 2)

    stats = {
        "at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "version": db["version"], "total": total,
        "source": "CS2 游戏自身 cvarlist 导出（VConsole2 会话日志，权威）"
                  " + ArmynC CS2 Cvars + OhLasFar release 清单",
        "fixed": len(fixed), "core_fixed": sum(1 for r in fixed if r["k"]),
        "old_english": old_english, "old_mixed": old_mixed, "with_zh": with_zh,
        "residual": residual, "hand": hand,
    }
    build_html(rows, stats, DOCS / "translation-report.html")
    (DOCS / "translation-report.json").write_text(
        json.dumps({"stats": stats, "rows": rows}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(stats, ensure_ascii=False, indent=1))

    docx_path = DOCS / "translation-report.docx"
    build_docx(rows, stats, docx_path, limit=limit)
    print(f"[OK] DOCX: {docx_path}  ({docx_path.stat().st_size/1024:.0f} KB)")
    try:
        to_pdf(docx_path, DOCS / "translation-report.pdf")
        print(f"[OK] PDF : {DOCS / 'translation-report.pdf'}")
    except Exception as e:
        print(f"[!] Word 导出 PDF 失败：{type(e).__name__} {e}")


if __name__ == "__main__":
    main()
